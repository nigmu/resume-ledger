"""Render a resume.md into .docx and .pdf using base/base_resume.docx.

Usage:
    python tools/render.py output/<job>/v1/resume.md [--preview] [--max-pages 1]

resume.md format (rules.md section 7 has the full spec):
    # Full Name
    +91 ... | [me@mail.com](mailto:me@mail.com) | [linkedin.com/in/me](https://linkedin.com/in/me)

    ## Education
    ### Bachelor of Technology | Example University | City, State
    #### Computer Science, **CGPA: 8.1** | 2021 – 2025

    ## Experience
    ### Company | City, Country | September 2025 – Present
    #### Job Title |
    - Bullet text.

    ## Skills
    **Languages:** Python, SQL

    ## Certificates
    [Certificate Name](https://example.com) – Issuer

`### A | B | R` prints "A," in bold, then B, then R right-aligned in the column.
`#### A | R` prints A in italics, then R upright in the column.
The last field is always the right-hand one; end with `|` to leave it empty.
Education's second line is indented. A job title uses the company indent.
Inline markup: **bold**, *italic*, [text](url).

The script writes <First_Last>_Resume.docx and .pdf next to resume.md, then
prints a fit report. Exit code 2 means a fit target was missed (page limit,
free lines outside 0-3, a long bullet, a short tail, or text running into
the right-hand column), not only a page overflow.
"""

import argparse
import copy
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Twips

sys.path.insert(0, str(Path(__file__).resolve().parent))
from make_base import (  # noqa: E402
    BODY_LINE, BULLET_LEFT, BULLET_RIGHT, COLUMN_TAB, ENTRY_GAP, LINE_GAP,
    MARGIN_BOTTOM, MARGIN_LEFT, MARGIN_RIGHT, MARGIN_TOP,
)

ROOT = Path(__file__).resolve().parent.parent
BASE = ROOT / "base" / "base_resume.docx"

INLINE = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*|\[([^\]]+)\]\(([^)\s]+)\)")
PLACEHOLDER = re.compile(r"\b(TODO|TBD|XXX|lorem ipsum)\b|\?\?|\{\{|\}\}", re.I)


def fail(msg):
    print(f"error: {msg}", file=sys.stderr)
    sys.exit(1)


def has_bad_markup(text):
    """True for placeholders, leftover '**', or a stray unclosed link.

    A single literal asterisk is allowed. Angle brackets are not placeholders.
    """
    stripped = INLINE.sub(lambda m: m.group(1) or m.group(2) or m.group(3), text)
    return bool(
        PLACEHOLDER.search(stripped)
        or "**" in stripped
        or "[" in stripped
        or "]" in stripped
    )


def version_path_error(md_path):
    """resume.md must be output/<job>/vN/resume.md, and vN must not be frozen."""
    output_root = (ROOT / "output").resolve()
    resolved = md_path.resolve()
    try:
        rel = resolved.relative_to(output_root)
    except ValueError:
        return f"refusing to render {resolved}: resume.md must be inside output/<job>/vN/"
    parts = rel.parts
    if len(parts) != 3 or parts[-1] != "resume.md" or not re.fullmatch(r"v\d+", parts[1]):
        return (
            f"refusing to render {resolved}: expected output/<job>/vN/resume.md, "
            "where vN is a version folder such as v1 or v2"
        )
    if (resolved.parent / "changes.md").is_file():
        return (
            f"refusing to render {resolved}: changes.md is already in this version, "
            "so the folder is frozen"
        )
    return None


def clear_previews(md_path):
    for png in md_path.parent.glob("preview_p*.png"):
        try:
            png.unlink()
        except OSError as exc:
            print(f"error: could not remove stale preview {png} ({exc})", file=sys.stderr)


def parse(md_text):
    """Return (name, contact_lines, blocks). blocks: list of (kind, payload)."""
    md_text = re.sub(r"<!--.*?-->", "", md_text, flags=re.S)
    name, contact, blocks = None, [], []
    for raw in md_text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            name = line[2:].strip()
        elif line.startswith("## "):
            blocks.append(("heading", line[3:].strip()))
        elif line.startswith("### "):
            blocks.append(("entry", [f.strip() for f in line[4:].split("|")]))
        elif line.startswith("#### "):
            blocks.append(("sub", [f.strip() for f in line[5:].split("|")]))
        elif line.startswith(("- ", "* ")):
            blocks.append(("bullet", line[2:].strip()))
        elif not blocks:
            contact.append(line)
        else:
            blocks.append(("text", line))
    if not name:
        fail("resume.md has no '# Name' line")
    return name, contact, blocks


def base_headings(doc):
    return [p.text.strip() for p in doc.paragraphs
            if p.style.style_id == "Heading1" and p.text.strip()]


def check_structure(name, contact, blocks, allowed):
    problems = []
    lowered = [h.lower() for h in allowed]
    last = -1
    for kind, h in blocks:
        if kind != "heading":
            continue
        if h.lower() not in lowered:
            problems.append(f"heading '{h}' is not in the base template ({', '.join(allowed)})")
            continue
        idx = lowered.index(h.lower())
        if idx <= last:
            problems.append(f"heading '{h}' is out of order; base order is {', '.join(allowed)}")
        last = idx
    if blocks and blocks[0][0] != "heading":
        problems.append("content appears before the first '## Section' heading")
    for i, (kind, _) in enumerate(blocks):
        if kind == "sub" and (i == 0 or blocks[i - 1][0] not in ("entry", "bullet", "sub")):
            problems.append("a '####' line must follow a '###' entry or its bullets")
    for payload in [name, *contact, *(b[1] for b in blocks)]:
        for t in (payload if isinstance(payload, list) else [payload]):
            if has_bad_markup(t):
                problems.append(f"possible placeholder or broken markup: {t}")
    return problems


def add_hyperlink(paragraph, text, url, bold=None, italic=None):
    if re.fullmatch(r"[^@\s:]+@[^@\s]+\.\w+", url):
        url = "mailto:" + url
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rstyle = OxmlElement("w:rStyle")
    rstyle.set(qn("w:val"), "Hyperlink")
    rpr.append(rstyle)
    for tag, value in (("w:b", bold), ("w:i", italic)):
        if value is not None:
            el = OxmlElement(tag)
            if not value:
                el.set(qn("w:val"), "0")
            rpr.append(el)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_run(paragraph, text, bold=None, italic=None):
    """None inherits from the paragraph style; True/False override it."""
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic


def add_inline(paragraph, text, bold=None, italic=None):
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            add_run(paragraph, text[pos:m.start()], bold, italic)
        if m.group(1) is not None:
            add_run(paragraph, m.group(1), True, italic)
        elif m.group(2) is not None:
            add_run(paragraph, m.group(2), bold, True)
        else:
            add_hyperlink(paragraph, m.group(3), m.group(4), bold, italic)
        pos = m.end()
    if pos < len(text):
        add_run(paragraph, text[pos:], bold, italic)


def split_fields(fields):
    """Return (first, middle_fields, right) using the 'last field is right' rule."""
    right = fields[-1] if len(fields) > 1 else ""
    middle = [f for f in fields[1:-1] if f]
    return fields[0], middle, right


def build_docx(name, contact, blocks, out_path):
    doc = Document(str(BASE))
    body = doc.element.body
    sect = copy.deepcopy(body.sectPr)
    for child in list(body):
        body.remove(child)
    body.append(sect)

    doc.add_paragraph(name, style="Name")
    for line in contact:
        add_inline(doc.add_paragraph(style="Contact"), line)

    kinds = []
    section = ""
    for kind, payload in blocks:
        if kind == "heading":
            section = payload.strip().lower()
            p = doc.add_paragraph(payload, style="heading 1")
            p.add_run("\t")
        elif kind == "entry":
            first, middle, right = split_fields(payload)
            p = doc.add_paragraph(style="Entry")
            add_inline(p, first + ("," if middle else ""), bold=True)
            if middle:
                add_inline(p, " " + ", ".join(middle))
            if right:
                p.add_run("\t")
                add_inline(p, right)
        elif kind == "sub":
            first, middle, right = split_fields(payload)
            # Education's field line sits further in. A job title lines up with the company.
            style_name = "Entry Sub" if section == "education" else "Entry Title"
            p = doc.add_paragraph(style=style_name)
            add_inline(p, ", ".join(f for f in [first, *middle] if f))
            if right:
                p.add_run("\t")
                add_inline(p, right, italic=False)
        elif kind == "bullet":
            p = doc.add_paragraph(style="Bullet")
            add_inline(p, payload)
        else:
            p = doc.add_paragraph(style="Normal")
            add_inline(p, payload)
        kinds.append((kind, section, p))

    # 2 pt under a bullet, a company line, and a project line.
    # 4 pt under the last line of a section, and before another degree.
    # A degree line and a job title sit flush on the line below them.
    for i, (kind, sec, p) in enumerate(kinds):
        next_kind = kinds[i + 1][0] if i + 1 < len(kinds) else "end"
        if kind == "heading":
            continue
        if next_kind in ("heading", "end") or (kind in ("sub", "text") and next_kind == "entry"):
            gap = ENTRY_GAP
        elif kind == "bullet" or (kind == "entry" and sec in ("experience", "projects")):
            gap = LINE_GAP
        else:
            gap = 0
        p.paragraph_format.space_after = Twips(gap)

    props = doc.core_properties
    props.author = name
    props.last_modified_by = name
    props.title = f"{name} - Resume"
    props.subject = props.keywords = props.comments = props.category = ""
    props.revision = 1
    doc.save(str(out_path))


def discard_failed_pdf(pdf_path):
    """Drop a pdf that would no longer match the docx. Say so if it stays locked."""
    if not pdf_path.is_file():
        return
    try:
        pdf_path.unlink()
    except OSError as exc:
        print(
            f"error: PDF export failed and the existing PDF is still here "
            f"because it could not be removed ({exc}): {pdf_path}",
            file=sys.stderr,
        )
    else:
        print(f"removed pdf left from the failed export: {pdf_path}", file=sys.stderr)


def libreoffice_executable():
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    for candidate in (
        "/usr/bin/soffice",
        "/usr/bin/libreoffice",
        "/usr/lib/libreoffice/program/soffice",
    ):
        if Path(candidate).is_file():
            return candidate
    return None


def export_with_libreoffice(docx_path, pdf_path):
    binary = libreoffice_executable()
    if binary is None:
        print(
            "error: LibreOffice is not installed; PDF export failed. "
            "Install it and make sure `soffice` is on PATH.",
            file=sys.stderr,
        )
        discard_failed_pdf(pdf_path)
        return False

    with tempfile.TemporaryDirectory(prefix="resume-lo-") as profile:
        command = [
            binary,
            f"-env:UserInstallation={Path(profile).resolve().as_uri()}",
            "--headless",
            "--norestore",
            "--convert-to",
            "pdf",
            "--outdir",
            str(pdf_path.parent),
            str(docx_path),
        ]
        try:
            result = subprocess.run(
                command, capture_output=True, text=True, timeout=120)
        except (OSError, subprocess.TimeoutExpired) as exc:
            print(f"error: LibreOffice export failed ({exc})", file=sys.stderr)
            discard_failed_pdf(pdf_path)
            return False

    produced = pdf_path.parent / f"{docx_path.stem}.pdf"
    if result.returncode != 0 or not produced.is_file():
        detail = (result.stderr or result.stdout or "no PDF was written").strip()
        print(f"error: LibreOffice export failed ({detail})", file=sys.stderr)
        discard_failed_pdf(pdf_path)
        return False
    if produced.resolve() != pdf_path.resolve():
        produced.replace(pdf_path)
    return True


def export_with_word(docx_path, pdf_path):
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        print("error: pywin32 is not installed; PDF export failed", file=sys.stderr)
        discard_failed_pdf(pdf_path)
        return False

    initialized = False
    word = None
    doc = None
    exported = False
    error = None
    try:
        pythoncom.CoInitialize()
        initialized = True
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(
            str(docx_path), ReadOnly=True, AddToRecentFiles=False, Visible=False)
        doc.ExportAsFixedFormat(
            str(pdf_path), 17, OptimizeFor=0, CreateBookmarks=1,
            DocStructureTags=True, BitmapMissingFonts=True)
        exported = True
    except Exception as exc:
        error = exc
    finally:
        if doc is not None:
            try:
                doc.Close(SaveChanges=0)
            except Exception as exc:
                print(f"warning: could not close Word ({exc})", file=sys.stderr)
        if word is not None:
            try:
                word.Quit()
            except Exception as exc:
                print(f"warning: could not quit Word ({exc})", file=sys.stderr)
        if initialized:
            pythoncom.CoUninitialize()
    if not exported:
        print(
            f"error: Word export failed ({error}). "
            "If the PDF is open in another program, close it and render again.",
            file=sys.stderr,
        )
        discard_failed_pdf(pdf_path)
        return False
    return True


def export_pdf(docx_path, pdf_path):
    """Word on Windows. LibreOffice everywhere else, including Linux."""
    if sys.platform == "win32":
        return export_with_word(docx_path, pdf_path)
    return export_with_libreoffice(docx_path, pdf_path)


def plain(text):
    return INLINE.sub(lambda m: m.group(1) or m.group(2) or m.group(3), text).strip()


def row_text(segs):
    return "".join(seg[2] for seg in segs)


def pdf_rows(page):
    """Visual text rows on a page: list of (y0, y1, [(x0, x1, text), ...]) top to bottom."""
    rows = []
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            text = "".join(s["text"] for s in line["spans"])
            if not text.strip():
                continue
            x0, y0, x1, y1 = line["bbox"]
            for row in rows:
                if abs(row[0] - y0) < 2:
                    row[2].append((x0, x1, text))
                    break
            else:
                rows.append([y0, y1, [(x0, x1, text)]])
    for row in rows:
        row[2].sort()
    return sorted(rows)


def bullet_limits(blocks):
    """Max lines for each bullet, in document order. Experience may run to 4."""
    section = ""
    limits = []
    for kind, payload in blocks:
        if kind == "heading":
            section = payload.strip().lower()
        elif kind == "bullet":
            limits.append(4 if section == "experience" else 2)
    return limits


def measure(pdf_path, blocks, preview):
    import pymupdf

    line_h = BODY_LINE / 20
    column_right = MARGIN_LEFT / 20 + COLUMN_TAB / 20
    bullet_x = MARGIN_LEFT / 20 + BULLET_LEFT / 20
    limits = bullet_limits(blocks)
    report = {"long_bullets": [], "short_tails": [], "bad_entries": []}
    with pymupdf.open(str(pdf_path)) as pdf:
        report["pages"] = len(pdf)
        page_h, page_w = pdf[0].rect.height, pdf[0].rect.width
        top, bottom = MARGIN_TOP / 20, page_h - MARGIN_BOTTOM / 20
        right_edge = page_w - MARGIN_RIGHT / 20 - BULLET_RIGHT / 20
        pages = [pdf_rows(p) for p in pdf]
        if preview:
            for i, page in enumerate(pdf):
                page.get_pixmap(dpi=110).save(str(pdf_path.with_name(f"preview_p{i + 1}.png")))

    used = [((rows[-1][1] - top) / line_h if rows else 0) for rows in pages]
    report["free_lines"] = round((bottom - top) / line_h - used[-1], 1)
    report["used_by_page"] = used

    rows = [r for page in pages for r in page]
    report["long_bullets"], report["short_tails"] = classify_bullets(
        rows, limits, bullet_x, right_edge)
    report["bad_entries"] = column_problems(blocks, rows, column_right)
    return report


def classify_bullets(rows, limits, bullet_x, right_edge):
    """Long-bullet and short-tail warnings.

    Word draws the bullet glyph on its own span, about 0.1 pt off the text,
    and pdf_rows merges those spans into one row. The warning text is every
    span on the row, and the cut estimate uses the last line's own text.
    """
    bullets, in_bullet = [], False
    for _, _, segs in rows:
        text = row_text(segs)
        x0 = segs[0][0]
        if text.lstrip().startswith("\u2022"):
            body = text.lstrip("\u2022 ").strip()
            bullets.append({
                "text": body,
                "last_line": body,
                "lines": 1,
                "end": segs[-1][1],
            })
            in_bullet = True
        elif in_bullet and abs(x0 - bullet_x) < 1.5:
            line = text.strip()
            bullets[-1]["lines"] += 1
            bullets[-1]["end"] = segs[-1][1]
            bullets[-1]["last_line"] = line
            if line:
                bullets[-1]["text"] = f"{bullets[-1]['text']} {line}".strip()
        else:
            in_bullet = False

    long_bullets, short_tails = [], []
    width = right_edge - bullet_x
    for i, b in enumerate(bullets):
        snippet = b["text"][:60] + ("..." if len(b["text"]) > 60 else "")
        fill = (b["end"] - bullet_x) / width if width else 0
        cap = limits[i] if i < len(limits) else 2
        if b["lines"] > cap:
            long_bullets.append((b["lines"], cap, snippet))
        elif b["lines"] >= 2 and fill < 0.3:
            chars = max(1, len(b["last_line"]))
            short_tails.append((round(fill * 100), chars, snippet))
    return long_bullets, short_tails


def column_problems(blocks, rows, column_right):
    """Flag an entry or subtitle whose text runs into the right-hand column.

    Lines with a date or location keep the old check. Lines with an empty
    right field, including project titles, are flagged when their ink reaches
    the column tab. Glyph boxes sit a few points inside that tab.
    """
    problems = []
    for kind, payload in blocks:
        if kind not in ("entry", "sub") or not payload:
            continue
        left_full = plain(payload[0])
        probe = left_full[:20]
        if not probe:
            continue
        has_right = len(payload) > 1 and bool(plain(payload[-1]))
        for _, _, segs in rows:
            head = segs[0][2].strip()
            joined = row_text(segs).strip()
            if not head.startswith(probe) and not joined.startswith(probe):
                continue
            if has_right:
                right_text = plain(payload[-1])
                on_row = [s for s in segs[1:] if s[2].strip().startswith(right_text[:10])]
                if not on_row or abs(on_row[0][1] - column_right) > 4:
                    problems.append(
                        f"text before the date/location column is too long: {left_full} ... {right_text}")
            else:
                rightmost = max(s[1] for s in segs)
                if rightmost > column_right - 12:
                    problems.append(f"text runs into the right-hand column: {left_full}")
            break
    return problems


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("resume_md", type=Path)
    ap.add_argument("--max-pages", type=int, default=1)
    ap.add_argument("--preview", action="store_true", help="also write preview_pN.png of each page")
    ap.add_argument("--no-pdf", action="store_true")
    args = ap.parse_args()

    md_path = args.resume_md.resolve()
    if not md_path.is_file():
        fail(f"{md_path} not found")
    if not BASE.is_file():
        fail(f"{BASE} not found; run tools/make_base.py")
    path_error = version_path_error(md_path)
    if path_error:
        fail(path_error)
    clear_previews(md_path)

    name, contact, blocks = parse(md_path.read_text(encoding="utf-8-sig"))
    problems = check_structure(name, contact, blocks, base_headings(Document(str(BASE))))
    if problems:
        for p in problems:
            print(f"error: {p}", file=sys.stderr)
        sys.exit(1)

    stem = re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_") + "_Resume"
    docx_path = md_path.with_name(stem + ".docx")
    pdf_path = md_path.with_name(stem + ".pdf")
    build_docx(name, contact, blocks, docx_path)
    print(f"docx: {docx_path}")
    if args.no_pdf:
        return

    if not export_pdf(docx_path, pdf_path):
        sys.exit(1)
    print(f"pdf:  {pdf_path}")
    report = measure(pdf_path, blocks, args.preview)
    over = report["pages"] > args.max_pages
    free = report["free_lines"]
    print(f"pages: {report['pages']} (limit {args.max_pages}){'  OVER LIMIT' if over else ''}")
    print(f"free lines at bottom: {free}")
    if over:
        extra = sum(report["used_by_page"][args.max_pages:])
        print(f"overflow: about {extra:.0f} lines past the limit (cut at least this much)")
    elif not 0 <= free <= 3:
        print(f"free lines miss: {free} is outside 0-3")
    for msg in report["bad_entries"]:
        print(msg)
    for lines, cap, snippet in report["long_bullets"]:
        print(f"bullet runs {lines} lines (max {cap}): {snippet}")
    for pct, chars, snippet in report["short_tails"]:
        print(f"bullet last line only {pct}% full (cut ~{chars} chars to save a line): {snippet}")
    free_miss = not over and not (0 <= free <= 3)
    if over or free_miss or report["long_bullets"] or report["short_tails"] or report["bad_entries"]:
        sys.exit(2)


if __name__ == "__main__":
    main()
